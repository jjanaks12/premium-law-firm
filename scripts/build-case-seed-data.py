from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
import uuid
from pathlib import Path

import openpyxl


NEPALI_DIGITS = str.maketrans("०१२३४५६७८९", "0123456789")
ZERO_WIDTH = re.compile(r"[\u200b-\u200f\u202a-\u202e\u2060\ufeff]")
SPACE = re.compile(r"\s+")
CASE_NUMBER_SEPARATOR = re.compile(r"\s*[,;।]+\s*")
NAMESPACE = uuid.UUID("8f486b80-f296-5d69-a8a9-99d37838e612")


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFC", str(value))
    text = ZERO_WIDTH.sub("", text)
    text = SPACE.sub(" ", text).strip(" ,")
    return text or None


def ascii_digits(value: str | None) -> str | None:
    return value.translate(NEPALI_DIGITS) if value else None


def normalize_case_number(value: str | None) -> str | None:
    value = ascii_digits(clean(value))
    if not value:
        return None
    parts = []
    for part in CASE_NUMBER_SEPARATOR.split(value):
        normalized = re.sub(r"\s*-\s*", "-", part).upper()
        normalized = re.sub(r"\s+", "", normalized)
        if normalized:
            parts.append(normalized)
    return ", ".join(parts) or None


def normalize_date(value: str | None) -> str | None:
    value = ascii_digits(clean(value))
    if not value:
        return None
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", value)
    if not match:
        return value
    year, month, day = (int(part) for part in match.groups())
    if not (2000 <= year <= 2199 and 1 <= month <= 12 and 1 <= day <= 32):
        return value
    return f"{year:04d}-{month:02d}-{day:02d}"


def comparable(value: str | None) -> str:
    value = ascii_digits(clean(value)) or ""
    return re.sub(r"[^\w\u0900-\u097f]+", "", value).lower()


def case_signature(row: dict) -> tuple[str, ...]:
    return tuple(
        comparable(row[field])
        for field in (
            "registrationDateBs",
            "firstParty",
            "matter",
            "court",
            "caseNumber",
            "remarks",
        )
    )


def merge_same_case_rows(rows: list[dict]) -> tuple[list[dict], int, int]:
    rows_by_signature: dict[tuple[str, ...], list[dict]] = {}
    for row in rows:
        rows_by_signature.setdefault(case_signature(row), []).append(row)

    merged: list[dict] = []
    cases_with_multiple_opposing_parties = 0
    for grouped_rows in rows_by_signature.values():
        latest = dict(grouped_rows[-1])
        latest["sourceRows"] = grouped_rows
        opposing_by_key: dict[str, str] = {}
        for row in grouped_rows:
            opposing = clean(row["opposingParty"])
            if opposing:
                opposing_by_key.setdefault(comparable(opposing), opposing)
        latest["opposingParties"] = list(opposing_by_key.values())
        if len(latest["opposingParties"]) > 1:
            cases_with_multiple_opposing_parties += 1
        merged.append(latest)

    merged.sort(key=lambda row: row["row"])
    return merged, len(rows) - len(merged), cases_with_multiple_opposing_parties


def deterministic_id(key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, key))


def nature_for(matter: str | None, case_number: str | None) -> str:
    haystack = f"{matter or ''} {case_number or ''}".lower()
    if re.search(r"(?:^|[-, ])(?:wo|wc|wp|writ)(?:[-, ]|$)", haystack) or any(
        word in haystack for word in ("उत्प्रेषण", "परमादेश", "निषेधाज्ञ", "बन्दीप्रत्यक्षीकरण")
    ):
        return "Writ"
    if re.search(r"(?:^|[-, ])(?:cr|c1)(?:[-, ]|$)", haystack) or any(
        word in haystack
        for word in ("ठगी", "जबरजस्ती", "मानव बेच", "लागु औषध", "चोरी", "किर्ते", "ज्यान", "फौजदारी")
    ):
        return "Criminal"
    if any(word in haystack for word in ("कम्पनी", "करार", "बैंक", "चेक अनादर", "व्यापार", "वाणिज्य")):
        return "Commercial"
    return "Civil"


def court_level_for(court: str | None) -> str | None:
    court = clean(court) or ""
    if re.search(r"सर्वो?च्च", court):
        return "Supreme Court"
    if re.search(r"उच्च(?:\s*अदालत)?", court):
        return "High Court"
    if re.search(r"जिल्ला(?:\s*अदालत)?|का\s*\.?\s*जि(?:\s*\.?\s*अ)?", court):
        return "District Court"
    if re.search(r"विशेष(?:\s*अदालत)?", court):
        return "Special Court"
    return None


def extract_rows(workbook_path: Path) -> tuple[list[dict], dict]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    expected = {"all cases", "running cases"}
    actual = set(workbook.sheetnames)
    if not expected.issubset(actual):
        raise ValueError(f"Expected sheets {sorted(expected)}, found {workbook.sheetnames}")

    rows: list[dict] = []
    skipped_blank = 0
    for sheet_name in ("all cases", "running cases"):
        sheet = workbook[sheet_name]
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = [clean(value) for value in cells[:8]]
            values.extend([None] * (8 - len(values)))
            if not any(values[1:8]):
                skipped_blank += 1
                continue
            serial, source_date, first_party, opposing_party, matter, court, case_number, remarks = values
            issues = []
            for label, value in (("first party", first_party), ("matter", matter), ("court", court)):
                if not value:
                    issues.append(f"Missing {label}")
            if not opposing_party:
                issues.append("Missing opposing party")
            normalized_number = normalize_case_number(case_number)
            if not normalized_number:
                issues.append("Missing case number")
            normalized_date = normalize_date(source_date)
            if not normalized_date:
                issues.append("Missing registration date")
            elif not re.fullmatch(r"\d{4}-\d{2}-\d{2}", normalized_date):
                issues.append("Unrecognized registration date")
            rows.append(
                {
                    "sheet": sheet_name,
                    "row": row_number,
                    "serial": ascii_digits(serial),
                    "registrationDateBs": normalized_date,
                    "firstParty": first_party,
                    "opposingParty": opposing_party,
                    "matter": matter,
                    "court": court,
                    "caseNumber": normalized_number,
                    "sourceCaseNumber": case_number,
                    "remarks": remarks,
                    "issues": issues,
                }
            )

    return rows, {"skippedBlankRows": skipped_blank, "sheetNames": workbook.sheetnames}


def merge_rows(rows: list[dict]) -> tuple[list[dict], dict]:
    dates_by_number: dict[str, set[str]] = {}
    for row in rows:
        date = row["registrationDateBs"]
        if not date or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date) or not row["caseNumber"]:
            continue
        for number in CASE_NUMBER_SEPARATOR.split(row["caseNumber"]):
            dates_by_number.setdefault(number, set()).add(date)

    propagated_dates = 0
    for row in rows:
        own_date = row["registrationDateBs"]
        if own_date and re.fullmatch(r"\d{4}-\d{2}-\d{2}", own_date):
            row["resolvedRegistrationDateBs"] = own_date
            continue
        candidates: set[str] = set()
        if row["caseNumber"]:
            for number in CASE_NUMBER_SEPARATOR.split(row["caseNumber"]):
                candidates.update(dates_by_number.get(number, set()))
        row["resolvedRegistrationDateBs"] = next(iter(candidates)) if len(candidates) == 1 else None
        if row["resolvedRegistrationDateBs"]:
            propagated_dates += 1

    all_source_rows = [row for row in rows if row["sheet"] == "all cases"]
    running_source_rows = [row for row in rows if row["sheet"] == "running cases"]
    all_rows, duplicate_all_rows_removed, _ = merge_same_case_rows(all_source_rows)
    running_rows, duplicate_running_rows_removed, _ = merge_same_case_rows(running_source_rows)

    by_number: dict[str, list[dict]] = {}
    by_composite: dict[str, list[dict]] = {}
    for row in all_rows:
        if row["caseNumber"]:
            for number in CASE_NUMBER_SEPARATOR.split(row["caseNumber"]):
                by_number.setdefault(number, []).append(row)
        composite = "|".join(comparable(row[field]) for field in ("firstParty", "matter", "court"))
        if composite.replace("|", ""):
            by_composite.setdefault(composite, []).append(row)

    matches: dict[int, dict] = {}
    ambiguous_matches = 0
    for running in running_rows:
        candidates: list[dict] = []
        if running["caseNumber"]:
            for number in CASE_NUMBER_SEPARATOR.split(running["caseNumber"]):
                candidates.extend(by_number.get(number, []))
        unique_candidates = {candidate["row"]: candidate for candidate in candidates}
        if len(unique_candidates) > 1:
            running_composite = "|".join(
                comparable(running[field]) for field in ("firstParty", "matter", "court")
            )
            matching_composite = [
                candidate
                for candidate in unique_candidates.values()
                if "|".join(
                    comparable(candidate[field])
                    for field in ("firstParty", "matter", "court")
                )
                == running_composite
            ]
            if matching_composite:
                candidates = matching_composite
        if not candidates:
            composite = "|".join(comparable(running[field]) for field in ("firstParty", "matter", "court"))
            candidates = by_composite.get(composite, [])
        unique_candidates = {candidate["row"]: candidate for candidate in candidates}
        if len(unique_candidates) == 1:
            matches[running["row"]] = next(iter(unique_candidates.values()))
        elif len(unique_candidates) > 1:
            ambiguous_matches += 1

    merged: list[dict] = []
    running_by_all_row: dict[int, list[dict]] = {}
    for running in running_rows:
        matched = matches.get(running["row"])
        if matched:
            running_by_all_row.setdefault(matched["row"], []).append(running)

    for row in all_rows:
        active_sources = running_by_all_row.get(row["row"], [])
        merged.append(build_case(row, active_sources, bool(active_sources)))

    matched_running_rows = set(matches)
    for row in running_rows:
        if row["row"] not in matched_running_rows:
            merged.append(build_case(row, [], True))

    duplicate_case_numbers = sum(len(group) - 1 for group in by_number.values() if len(group) > 1)
    report = {
        "sourceRows": len(rows),
        "allCaseRows": len(all_source_rows),
        "runningCaseRows": len(running_source_rows),
        "deduplicatedAllCaseRows": len(all_rows),
        "deduplicatedRunningCaseRows": len(running_rows),
        "matchedRunningRows": len(matches),
        "unmatchedRunningRows": len(running_rows) - len(matches),
        "ambiguousRunningMatches": ambiguous_matches,
        "generatedCases": len(merged),
        "activeCases": sum(case["status"] == "Active" for case in merged),
        "closedCases": sum(case["status"] == "Closed" for case in merged),
        "duplicateCaseNumbersInAllCases": duplicate_case_numbers,
        "sameCaseAllRowsMerged": duplicate_all_rows_removed,
        "sameCaseRunningRowsMerged": duplicate_running_rows_removed,
        "casesWithMultipleOpposingParties": sum(len(case["parties"]) > 2 for case in merged),
        "registrationDatesPropagatedByCaseNumber": propagated_dates,
        "rowsWithIssues": sum(bool(row["issues"]) for row in rows),
        "issueCounts": issue_counts(rows),
    }
    return merged, report


def build_case(primary: dict, active_sources: list[dict], active: bool) -> dict:
    source_key = f"{primary['sheet']}:{primary['row']}"
    case_id = deterministic_id(f"case:{source_key}")
    first_party_role = "Petitioner" if nature_for(primary["matter"], primary["caseNumber"]) == "Writ" else "Plaintiff"
    second_party_role = "Respondent" if first_party_role == "Petitioner" else "Defendant"
    source_rows = [
        *primary["sourceRows"],
        *(source for active_source in active_sources for source in active_source["sourceRows"]),
    ]
    opposing_by_key: dict[str, str] = {}
    for row in [primary, *active_sources]:
        for opposing in row["opposingParties"]:
            opposing_by_key.setdefault(comparable(opposing), opposing)
    opposing_parties = list(opposing_by_key.values()) or ["नाम उपलब्ध छैन"]
    registration_date_bs = next(
        (row["resolvedRegistrationDateBs"] for row in source_rows if row["resolvedRegistrationDateBs"]),
        None,
    )
    return {
        "id": case_id,
        "nature": nature_for(primary["matter"], primary["caseNumber"]),
        "status": "Active" if active else "Closed",
        "facts": primary["matter"],
        "details": {
            "importSource": "Copy of CASES LIST.xlsx",
            "sourceRows": [
                {
                    "sheet": row["sheet"],
                    "row": row["row"],
                    "serial": row["serial"],
                    "registrationDateBs": row["registrationDateBs"],
                    "sourceCaseNumber": row["sourceCaseNumber"],
                    "remarks": row["remarks"],
                    "validationIssues": row["issues"],
                }
                for row in source_rows
            ],
        },
        "courtDetail": {
            "id": deterministic_id(f"court-detail:{source_key}"),
            "caseName": " विरुद्ध ".join(filter(None, (primary["firstParty"], " / ".join(opposing_parties))))
            or primary["matter"]
            or f"Imported case {primary['row']}",
            "caseNumber": primary["caseNumber"] or "",
            "registrationDate": None,
            "registrationDateBs": registration_date_bs,
            "courtLevel": court_level_for(primary["court"]),
            "courtName": primary["court"],
            "isActive": active,
        },
        "parties": [
            {
                "id": deterministic_id(f"party:first:{source_key}"),
                "name": primary["firstParty"] or "नाम उपलब्ध छैन",
                "role": first_party_role,
            },
            *[
              {
                "id": deterministic_id(
                    f"party:opposing:{source_key}" if index == 0 else f"party:opposing:{index}:{source_key}"
                ),
                "name": opposing_party,
                "role": second_party_role,
              }
              for index, opposing_party in enumerate(opposing_parties)
            ],
        ],
    }


def issue_counts(rows: list[dict]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for row in rows:
        for issue in row["issues"]:
            counts[issue] = counts.get(issue, 0) + 1
    return dict(sorted(counts.items()))


def render_typescript(cases: list[dict], metadata: dict) -> str:
    payload = json.dumps(cases, ensure_ascii=False, indent=2)
    meta = json.dumps(metadata, ensure_ascii=False, indent=2)
    return (
        "// Generated by scripts/build-case-seed-data.py. Do not edit by hand.\n"
        "// Valid Bikram Sambat dates are stored in CaseCourtDetail.registrationDateBs.\n\n"
        "export const importedCaseSeedMetadata = "
        + meta
        + " as const;\n\nexport const importedCaseSeedData = "
        + payload
        + " as const;\n"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate deterministic Prisma case seed data from the case workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    source_bytes = args.workbook.read_bytes()
    rows, extraction = extract_rows(args.workbook)
    cases, report = merge_rows(rows)
    metadata = {
        "sourceFile": args.workbook.name,
        "sourceSha256": hashlib.sha256(source_bytes).hexdigest(),
        **extraction,
        **report,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(render_typescript(cases, metadata), encoding="utf-8", newline="\n")
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
